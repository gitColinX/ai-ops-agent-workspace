"""Build a blank intake workbook for agent-assisted onboarding.

The workbook turns a request form into a labeled text block that you paste
into an AI assistant instead of a screenshot. Layout on every entry tab:

    column A  label (bold when the field is required)
    column B  value (shaded input cell)
    column C  hidden row marker: "section", "field", "required", "local"
              or "output"

The "Copy for agent" macro (IntakeMacros.bas) and intake_paste_block.py both
read the markers in column C, so rows can be added or reordered without
changing any code. Rows marked "local" (personal email, mobile phone) stay
in the workbook for the person doing the build and are never copied into
the block, so they do not end up in an AI chat history.

openpyxl cannot create a VBA project, so this script writes an .xlsx. To get
the button: open the file in Excel, press Alt+F11, File > Import File,
choose IntakeMacros.bas, save as .xlsm, then run AddIntakeButtons once.

Usage:
    python -I build_intake_workbook.py --out Intake_Template.xlsx
    python -I build_intake_workbook.py --out Intake_Template.xlsx --overwrite

Writes one local file. Refuses to overwrite an existing file unless
--overwrite is given. Contacts nothing.
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

INPUT_FILL = PatternFill("solid", start_color="FFF6D5")
SECTION_FILL = PatternFill("solid", start_color="D9E2F3")
OUTPUT_FILL = PatternFill("solid", start_color="E2EFDA")

YES_NO = ["Yes", "No"]
OFFICES = ["Springfield HQ", "Riverside Field Office", "Remote"]
RESOURCE_TYPES = [
    "Contractor (staffing agency)",
    "Independent contractor",
    "Consultant (professional services firm)",
    "Service provider",
]

# Each row: (kind, label, hint, choices)
#   kind: "section", "field", "required" or "local" (kept out of the block)
#   choices: a list for a drop-down, "date" for a date check, or None
FORMS: dict[str, dict] = {
    "New Employee": {
        "title": "INTAKE: NEW EMPLOYEE",
        "rows": [
            ("section", "TICKET", None, None),
            ("required", "Ticket", "Ticket or task number from the service desk.", None),
            ("field", "Request item", "Request item number, if different from the ticket.", None),
            ("field", "Requested by", "Who submitted the request (UPN).", None),
            ("section", "PERSON", None, None),
            ("required", "Legal name", "First Last, exactly as HR has it.", None),
            ("field", "Preferred name", "Only if the display name or email should differ.", None),
            ("required", "Start date", "First day. Enter as a date.", "date"),
            ("required", "Manager", "Manager UPN, for example manager@contoso.com.", None),
            ("required", "Department", "Department name as it appears in the directory.", None),
            ("required", "Job title", "Title as HR has it.", None),
            ("field", "Employee ID", "HR employee ID, if the ticket has one.", None),
            ("local", "Personal email", "First-day contact only. Stays in this workbook; not copied for the agent.", None),
            ("local", "Mobile phone", "For MFA pre-registration, if policy allows. Stays in this workbook; not copied for the agent.", None),
            ("section", "PLACEMENT", None, None),
            ("required", "Office", "Home office. Drives the address fields.", OFFICES),
            ("field", "Intern", "Yes or No.", YES_NO),
            ("section", "HARDWARE AND SOFTWARE", None, None),
            ("required", "Laptop needed", "Yes or No.", YES_NO),
            ("field", "Delivery or pickup", "Where the laptop goes, if needed.", None),
            ("field", "Applications", "Applications requested, comma separated.", None),
            ("field", "Notes", "Anything else from the ticket.", None),
        ],
    },
    "Contractor": {
        "title": "INTAKE: CONTRACTOR OR CONSULTANT",
        "rows": [
            ("section", "TICKET", None, None),
            ("required", "Ticket", "Ticket or task number from the service desk.", None),
            ("field", "Request item", "Request item number, if different from the ticket.", None),
            ("required", "Approval record", "Number of the approval that authorizes this account.", None),
            ("section", "PERSON", None, None),
            ("required", "Name", "First Last.", None),
            ("required", "Resource type", "See the Field Notes tab.", RESOURCE_TYPES),
            ("required", "Resource company", "The agency or firm.", None),
            ("field", "Business email", "Agency or firm email address.", None),
            ("local", "Phone", "For MFA pre-registration, if policy allows. Stays in this workbook; not copied for the agent.", None),
            ("required", "Sponsor", "Internal manager or sponsor (UPN).", None),
            ("required", "Department", "Department the work is for.", None),
            ("section", "DATES", None, None),
            ("required", "Start date", "First day. Enter as a date.", "date"),
            ("required", "End date", "Estimated end. Becomes the account expiry.", "date"),
            ("section", "PLACEMENT", None, None),
            ("required", "Reporting office", "Office the person reports to.", OFFICES),
            ("field", "International user", "Yes or No. May change sign-in policy.", YES_NO),
            ("section", "HARDWARE AND SOFTWARE", None, None),
            ("required", "Laptop needed", "Yes or No.", YES_NO),
            ("field", "Delivery or pickup", "Where the laptop goes, if needed.", None),
            ("field", "Applications", "Applications requested, comma separated.", None),
            ("field", "Notes", "Anything else from the ticket.", None),
        ],
    },
}

FIELD_NOTES = [
    ("Contractor (staffing agency)", "Employed and paid by a staffing agency that places them here."),
    ("Independent contractor", "Engaged directly for a defined piece of work; not on payroll."),
    ("Consultant (professional services firm)", "Employee of a firm with a services agreement."),
    ("Service provider", "An organization contracted for a service; access is for its staff."),
    ("Employees", "New employees use the New Employee tab and have no resource type."),
]

FIRST_ROW = 4


def _add_entry_sheet(wb: Workbook, name: str, spec: dict) -> None:
    ws = wb.create_sheet(name)
    ws["A1"] = spec["title"]
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = (
        "Fill the shaded cells in column B. Hover over a label for guidance. "
        "Then click Copy for agent and paste the block into your agent session."
    )
    ws["A2"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A2:B2")
    ws.row_dimensions[2].height = 32

    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 58
    ws.column_dimensions["C"].hidden = True
    ws.column_dimensions["D"].width = 26

    validations: dict[str, DataValidation] = {}
    row = FIRST_ROW
    for kind, label, hint, choices in spec["rows"]:
        a = ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=3, value=kind)
        if kind == "section":
            a.font = Font(bold=True)
            a.fill = SECTION_FILL
            ws.cell(row=row, column=2).fill = SECTION_FILL
        else:
            a.font = Font(bold=(kind == "required"), italic=(kind == "local"))
            b = ws.cell(row=row, column=2)
            b.fill = INPUT_FILL
            b.alignment = Alignment(wrap_text=True, vertical="top")
            if hint:
                note = hint + (" Required." if kind == "required" else "")
                a.comment = Comment(note, "intake")
            if choices == "date":
                key = "date"
                if key not in validations:
                    dv = DataValidation(type="date", operator="greaterThan", formula1="DATE(2000,1,1)")
                    dv.error = "Enter a date."
                    dv.showErrorMessage = True
                    validations[key] = dv
                validations[key].add(b.coordinate)
                b.number_format = "yyyy-mm-dd"
            elif isinstance(choices, list):
                key = "|".join(choices)
                if key not in validations:
                    dv = DataValidation(type="list", formula1='"' + ",".join(choices) + '"', allow_blank=True)
                    dv.error = "Pick a value from the list."
                    dv.showErrorMessage = True
                    validations[key] = dv
                validations[key].add(b.coordinate)
        row += 1

    row += 1
    out = ws.cell(row=row, column=1, value="PASTE OUTPUT")
    out.font = Font(bold=True)
    out.fill = OUTPUT_FILL
    ws.cell(row=row, column=2, value="Fallback area: the macro writes the block here if the clipboard copy cannot be confirmed.")
    ws.cell(row=row, column=2).fill = OUTPUT_FILL
    ws.cell(row=row, column=3, value="output")

    for dv in validations.values():
        ws.add_data_validation(dv)
    ws.freeze_panes = "A4"


def _add_notes_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("Field Notes")
    ws["A1"] = "FIELD NOTES (reference, not an intake tab)"
    ws["A1"].font = Font(bold=True, size=14)
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 70
    for i, (term, meaning) in enumerate(FIELD_NOTES, start=3):
        ws.cell(row=i, column=1, value=term).font = Font(bold=True)
        ws.cell(row=i, column=2, value=meaning).alignment = Alignment(wrap_text=True)


def build_workbook(out_path: Path, overwrite: bool = False) -> Path:
    """Write the blank intake workbook to out_path and return the path."""
    out_path = Path(out_path)
    if out_path.suffix.lower() != ".xlsx":
        raise ValueError("Output must be an .xlsx file (add the macro in Excel, then save as .xlsm).")
    if out_path.exists() and not overwrite:
        raise FileExistsError(f"{out_path} exists. Use --overwrite to replace it.")
    wb = Workbook()
    wb.remove(wb.active)
    for name, spec in FORMS.items():
        _add_entry_sheet(wb, name, spec)
    _add_notes_sheet(wb)
    wb.active = 0
    wb.properties.creator = "build_intake_workbook.py"
    wb.properties.title = "Intake template"
    wb.save(out_path)
    return out_path


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--out", default="Intake_Template.xlsx", help="output .xlsx path")
    parser.add_argument("--overwrite", action="store_true", help="replace an existing file")
    args = parser.parse_args(argv)
    try:
        path = build_workbook(Path(args.out), overwrite=args.overwrite)
    except (FileExistsError, ValueError) as exc:
        print(f"Not written: {exc}", file=sys.stderr)
        return 1
    print(f"Wrote {path.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
