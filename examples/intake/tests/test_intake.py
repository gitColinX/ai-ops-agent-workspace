"""Tests for the intake generator and paste-block reader, on synthetic data.

Run from anywhere:
    python -I examples/intake/tests/test_intake.py

Everything is written to a temporary folder and removed afterwards.
"""

from __future__ import annotations

import datetime as dt
import importlib.util
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

sys.dont_write_bytecode = True  # leave no __pycache__ beside the examples

HERE = Path(__file__).resolve().parent
INTAKE_DIR = HERE.parent


def _load(name: str):
    spec = importlib.util.spec_from_file_location(name, INTAKE_DIR / f"{name}.py")
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


builder = _load("build_intake_workbook")
reader = _load("intake_paste_block")


def _row_of(ws, label: str) -> int:
    for row in range(1, ws.max_row + 1):
        if ws.cell(row=row, column=1).value == label and ws.cell(row=row, column=3).value in ("field", "required", "local"):
            return row
    raise AssertionError(f"label {label!r} not found on {ws.title}")


class IntakeTests(unittest.TestCase):
    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp = Path(self._tmp.name)
        self.blank = builder.build_workbook(self.tmp / "Intake_Template.xlsx")

    def tearDown(self) -> None:
        self._tmp.cleanup()

    def test_structure(self) -> None:
        wb = load_workbook(self.blank)
        self.assertEqual(wb.sheetnames, ["New Employee", "Contractor", "Field Notes"])
        for name in ("New Employee", "Contractor"):
            ws = wb[name]
            self.assertTrue(ws.column_dimensions["C"].hidden, f"{name}: column C should be hidden")
            markers = [ws.cell(row=r, column=3).value for r in range(1, ws.max_row + 1)]
            self.assertIn("required", markers)
            self.assertIn("local", markers)
            self.assertEqual(markers.count("output"), 1)
            self.assertGreater(len(ws.data_validations.dataValidation), 0)
            commented = [ws.cell(row=r, column=1).comment for r in range(1, ws.max_row + 1)]
            self.assertTrue(any(commented), f"{name}: labels should carry guidance comments")
        wb.close()

    def test_refuses_overwrite(self) -> None:
        with self.assertRaises(FileExistsError):
            builder.build_workbook(self.blank)
        self.assertEqual(builder.main(["--out", str(self.blank)]), 1)
        self.assertEqual(builder.main(["--out", str(self.blank), "--overwrite"]), 0)

    def test_rejects_xlsm_output(self) -> None:
        with self.assertRaises(ValueError):
            builder.build_workbook(self.tmp / "x.xlsm")

    def test_blank_form_lists_every_required_field(self) -> None:
        block = reader.paste_block(self.blank, "New Employee")
        required = [label for kind, label, _h, _c in builder.FORMS["New Employee"]["rows"] if kind == "required"]
        self.assertIn("--- Blank required fields: " + ", ".join(required), block)

    def test_filled_contractor_block(self) -> None:
        wb = load_workbook(self.blank)
        ws = wb["Contractor"]
        values = {
            "Ticket": "TASK-0001",
            "Approval record": "APR-0042",
            "Name": "Avery Sample",
            "Resource type": "Contractor (staffing agency)",
            "Resource company": "Fabrikam Staffing",
            "Business email": "avery@fabrikam.com",
            "Sponsor": "morgan.example@contoso.com",
            "Department": "Operations",
            "Start date": dt.datetime(2026, 1, 12),
            "End date": dt.datetime(2026, 6, 30),
            "Reporting office": "Springfield HQ",
            "Laptop needed": "Yes",
            "Notes": "Line one\nline two",
        }
        for label, value in values.items():
            ws.cell(row=_row_of(ws, label), column=2, value=value)
        filled = self.tmp / "filled.xlsx"
        wb.save(filled)
        wb.close()

        block = reader.paste_block(filled, "Contractor")
        lines = block.splitlines()
        self.assertEqual(lines[0], "=== INTAKE: CONTRACTOR OR CONSULTANT ===")
        self.assertEqual(lines[-1], "=== END INTAKE ===")
        self.assertIn("Ticket: TASK-0001", lines)
        self.assertIn("Start date: 2026-01-12", lines)
        self.assertIn("End date: 2026-06-30", lines)
        self.assertIn("Notes: Line one line two", lines)
        self.assertIn("International user: ", lines)
        self.assertFalse(any(line.startswith("Phone") for line in lines), "local rows must not reach the block")
        self.assertIn("--- Blank required fields: none", lines)
        self.assertFalse(any(line.startswith(("TICKET", "PERSON", "PASTE OUTPUT")) for line in lines))

    def test_one_required_blank(self) -> None:
        wb = load_workbook(self.blank)
        ws = wb["New Employee"]
        for kind, label, _h, choices in builder.FORMS["New Employee"]["rows"]:
            if kind == "required" and label != "Manager":
                value = dt.datetime(2026, 1, 12) if choices == "date" else "x"
                ws.cell(row=_row_of(ws, label), column=2, value=value)
        filled = self.tmp / "filled2.xlsx"
        wb.save(filled)
        wb.close()
        self.assertIn("--- Blank required fields: Manager", reader.paste_block(filled, "New Employee"))

    def test_local_rows_stay_out_of_the_block(self) -> None:
        wb = load_workbook(self.blank)
        ws = wb["New Employee"]
        ws.cell(row=_row_of(ws, "Personal email"), column=2, value="avery.home@example.com")
        ws.cell(row=_row_of(ws, "Mobile phone"), column=2, value="(555) 555-0142")
        filled = self.tmp / "filled3.xlsx"
        wb.save(filled)
        wb.close()
        block = reader.paste_block(filled, "New Employee")
        self.assertNotIn("avery.home@example.com", block)
        self.assertNotIn("555-0142", block)
        self.assertNotIn("Personal email", block)
        self.assertNotIn("Mobile phone", block)

    def test_reference_tab_has_no_block(self) -> None:
        self.assertEqual(reader.paste_block(self.blank, "Field Notes"), "")
        self.assertEqual(reader.main([str(self.blank), "--sheet", "Field Notes"]), 1)
        self.assertEqual(reader.main([str(self.blank), "--sheet", "Nope"]), 1)


if __name__ == "__main__":
    unittest.main(verbosity=2)
