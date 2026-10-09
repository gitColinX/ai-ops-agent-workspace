"""Print the paste block from a filled intake workbook.

Same output as the "Copy for agent" macro in IntakeMacros.bas, for people
who would rather not enable macros. Reads column A (label), column B
(value) and hidden column C (row marker) on one tab, stops at the
"output" marker, and lists required fields that are still blank so the
assistant asks about them instead of guessing. Rows marked "local" hold
personal contact details for the person doing the build and are never
printed.

Usage:
    python -I intake_paste_block.py Filled_Intake.xlsx
    python -I intake_paste_block.py Filled_Intake.xlsx --sheet Contractor

Read-only. Prints to stdout. Contacts nothing.
"""

from __future__ import annotations

import argparse
import datetime as dt
import sys
from pathlib import Path

from openpyxl import load_workbook

FIELD_MARKERS = {"field", "required"}  # "local" rows are deliberately left out


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, (dt.datetime, dt.date)):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return " ".join(str(value).replace("\r", " ").replace("\n", " ").split())


def build_block(ws) -> str:
    """Return the paste block for one worksheet, or "" if it has no intake rows."""
    lines = [f"=== {_cell_text(ws.cell(row=1, column=1).value)} ==="]
    blanks: list[str] = []
    found = False
    for row in range(1, ws.max_row + 1):
        marker = _cell_text(ws.cell(row=row, column=3).value).lower()
        if marker == "output":
            break
        if marker not in FIELD_MARKERS:
            continue
        found = True
        label = _cell_text(ws.cell(row=row, column=1).value)
        value = _cell_text(ws.cell(row=row, column=2).value)
        lines.append(f"{label}: {value}")
        if marker == "required" and not value:
            blanks.append(label)
    if not found:
        return ""
    lines.append("--- Blank required fields: " + (", ".join(blanks) if blanks else "none"))
    lines.append("=== END INTAKE ===")
    return "\n".join(lines)


def paste_block(path: Path, sheet: str | None = None) -> str:
    wb = load_workbook(Path(path), data_only=True)
    try:
        ws = wb[sheet] if sheet else wb.active
        return build_block(ws)
    finally:
        wb.close()


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("workbook", help="filled .xlsx or .xlsm")
    parser.add_argument("--sheet", help="tab name (default: the active tab)")
    args = parser.parse_args(argv)
    try:
        block = paste_block(Path(args.workbook), args.sheet)
    except KeyError:
        print(f"No tab named {args.sheet!r}.", file=sys.stderr)
        return 1
    if not block:
        print("That tab has no intake rows (hidden column C is empty).", file=sys.stderr)
        return 1
    print(block)
    return 0


if __name__ == "__main__":
    sys.exit(main())
